"""전체.xlsx — 시트 = 지역. 상태 열을 편집해 다시 불러올 수 있음 (run.py import-status)."""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import Finding

HEADERS = ["번호", "매체", "URL", "특정 병원", "위반 유형", "확신도", "근거", "캡처 파일", "상태", "신고일", "처리 결과", "키워드", "작성자"]
WIDTHS = [6, 8, 55, 22, 22, 8, 60, 24, 10, 12, 14, 14, 16]


def write_xlsx(path: Path, items: list[Finding], excluded: list[Finding] | None = None) -> Path:
    by_region: dict[str, list[Finding]] = defaultdict(list)
    for f in items:
        by_region[f.region].append(f)
    wb = Workbook()
    wb.remove(wb.active)
    if not by_region:
        ws = wb.create_sheet("없음")
        ws.append(["신고 후보 없음"])
    for region, group in by_region.items():
        ws = wb.create_sheet(region[:31])
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for f in sorted(group, key=lambda f: (f.medium, f.seq)):
            ws.append([f.seq, f.medium_label, f.url, f.clinic_name, f.violation_type, f.confidence, f.summary,
                       Path(f.screenshot).name if f.screenshot else "", f.status, f.reported_at, f.result,
                       f.keyword, f.author])
        for i, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            row[6].alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    if excluded:
        ws = wb.create_sheet("제외")
        ws.append(["지역", "매체", "URL", "상태", "사유", "특정 병원"])
        for f in excluded:
            ws.append([f.region, f.medium_label, f.url, f.status, f.summary, f.clinic_name])
        for i, w in enumerate([10, 8, 55, 10, 60, 22], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def read_status(path: Path) -> list[dict]:
    """[{url, status, reported_at, result}] — 사용자가 엑셀에서 편집한 상태를 읽어옴."""
    wb = load_workbook(path, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not {"URL", "상태", "신고일", "처리 결과"} <= set(header):
            continue  # "제외" 시트 등은 건너뜀
        idx = {h: i for i, h in enumerate(header)}
        for r in rows:
            url = r[idx["URL"]] if idx["URL"] < len(r) else None
            if not url:
                continue
            out.append({
                "url": str(url).strip(),
                "status": str(r[idx["상태"]] or "").strip(),
                "reported_at": _date_str(r[idx["신고일"]]),
                "result": str(r[idx["처리 결과"]] or "").strip(),
            })
    return out


def _date_str(v) -> str:
    if v is None or v == "":
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()
