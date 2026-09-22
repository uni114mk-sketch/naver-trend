"""신고기록_YYYY-MM.pptx — 내부 기록용. 신고함 상태(및 그 이후 상태)인 건만."""
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

from ..models import Finding

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)


def _title_only(prs, title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(28)
    return slide


def _add_table(slide, rows: list[list[str]], col_widths: list[float], top=1.3, font_pt=11):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.4), Inches(top), Inches(sum(col_widths)), Inches(0.4 * n_rows))
    table = shape.table
    for j, w in enumerate(col_widths):
        table.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font_pt if i else font_pt + 1)
                p.font.bold = i == 0
    return table


def write_pptx(path: Path, month: str, items: list[Finding], per_item_slides: bool = True) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    by_region: dict[str, list[Finding]] = defaultdict(list)
    for f in items:
        by_region[f.region].append(f)

    # 1. 표지
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = f"불법 의료광고 신고 기록 {month}"
    counts = Counter(f.region for f in items)
    sub = f"총 {len(items)}건 / 지역 {len(counts)}곳"
    if counts:
        sub += "\n" + ", ".join(f"{r} {n}" for r, n in sorted(counts.items(), key=lambda x: -x[1]))
    cover.placeholders[1].text = sub

    # 2. 지역별 요약
    for region in sorted(by_region):
        group = sorted(by_region[region], key=lambda f: (f.reported_at, f.medium, f.seq))
        rows = [["번호", "매체", "병원명", "위반 유형", "신고일", "결과", "URL"]]
        for i, f in enumerate(group, 1):
            rows.append([i, f.medium_label, f.clinic_name, f.violation_type, f.reported_at, f.result or f.status, f.url])
        for chunk_start in range(0, len(rows) - 1, 12):  # 슬라이드당 12행
            chunk = [rows[0]] + rows[1 + chunk_start: 1 + chunk_start + 12]
            slide = _title_only(prs, f"{region} — 신고 {len(group)}건")
            _add_table(slide, chunk, [0.6, 0.9, 2.0, 2.2, 1.1, 1.1, 4.6], font_pt=10)

    # 3. 건별 상세 (캡처 + 정보)
    if per_item_slides:
        for region in sorted(by_region):
            for f in sorted(by_region[region], key=lambda f: (f.reported_at, f.medium, f.seq)):
                slide = _title_only(prs, f"{region} · {f.medium_label} · {f.clinic_name or '병원 미상'}")
                shot = Path(f.screenshot) if f.screenshot else None
                if shot and shot.exists():
                    _add_picture_fit(slide, shot, Inches(0.4), Inches(1.3), Inches(5.6), Inches(5.8))
                else:
                    tb = slide.shapes.add_textbox(Inches(0.4), Inches(1.3), Inches(5.6), Inches(1))
                    tb.text_frame.text = "(캡처 없음)"
                info = slide.shapes.add_textbox(Inches(6.3), Inches(1.3), Inches(6.6), Inches(5.8))
                tf = info.text_frame
                tf.word_wrap = True
                lines = [
                    ("URL", f.url), ("특정 병원", f.clinic_name), ("위반 유형", f.violation_type),
                    ("근거", f.summary), ("신고일", f.reported_at), ("처리 결과", f.result or f.status),
                    ("검색 키워드", f.keyword), ("수집일", f.run_date),
                ]
                first = True
                for k, v in lines:
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    r1 = p.add_run()
                    r1.text = f"{k}: "
                    r1.font.bold = True
                    r1.font.size = Pt(13)
                    r2 = p.add_run()
                    r2.text = str(v or "-")
                    r2.font.size = Pt(13)
                    p.space_after = Pt(6)

    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def _add_picture_fit(slide, img_path: Path, left, top, max_w, max_h):
    from PIL import Image
    with Image.open(img_path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    slide.shapes.add_picture(str(img_path), left, top, width=Emu(int(w * scale)), height=Emu(int(h * scale)))
